from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('.')
OUT=ROOT/'reports/2026/2026-09-08_2026-09-14'; OUT.mkdir(parents=True,exist_ok=True)
sources=json.loads((ROOT/'.github/data/backfill_sources.json').read_text(encoding='utf-8'))
cases=[]
for _n in ('a','b','c'):
    cases += json.loads((ROOT/f'.github/data/backfill_cases_{_n}.json').read_text(encoding='utf-8'))
uncertain=json.loads((ROOT/'.github/data/backfill_uncertain.json').read_text(encoding='utf-8'))
LANGS=json.loads((ROOT/'.github/data/backfill_LANGS.json').read_text(encoding='utf-8'))
ACCENT='0A6B55'; DARK='14372C'; GRAY='5C6B66'

def shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr(); n=pr.find(qn('w:shd'))
    if n is None: n=OxmlElement('w:shd'); pr.append(n)
    n.set(qn('w:fill'),fill)

def margins(cell,v=80):
    pr=cell._tc.get_or_add_tcPr(); mar=pr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); pr.append(mar)
    for k in ('top','start','bottom','end'):
        n=mar.find(qn('w:'+k))
        if n is None: n=OxmlElement('w:'+k); mar.append(n)
        n.set(qn('w:w'),str(v)); n.set(qn('w:type'),'dxa')

def no_split(row):
    pr=row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))

def link(p,text,url):
    rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rid); rr=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    c=OxmlElement('w:color'); c.set(qn('w:val'),ACCENT); rp.append(c); u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); rp.append(u)
    rr.append(rp); t=OxmlElement('w:t'); t.text=text; rr.append(t); h.append(rr); p._p.append(h)

def clear_body(doc):
    body=doc._element.body
    for ch in list(body):
        if ch.tag != qn('w:sectPr'): body.remove(ch)

def body(doc,text,size=10,after=6,italic=False):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.08
    r=p.add_run(text); r.font.name='Aptos'; r.font.size=Pt(size); r.italic=italic
    return p

def heading(doc,text,level=1):
    p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(10 if level==1 else 7); p.paragraph_format.space_after=Pt(5)
    r=p.add_run(text); r.bold=True; r.font.name='Aptos'; r.font.size=Pt(16 if level==1 else 11.5); r.font.color.rgb=RGBColor.from_string(DARK if level==1 else ACCENT)

def val(c,lang,k):
    if lang=='pt': return c[k]
    return c.get(k+'_'+lang,c[k])

def first_page(doc,cfg):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); r=p.add_run(cfg['title']); r.bold=True; r.font.name='Aptos'; r.font.size=Pt(18); r.font.color.rgb=RGBColor.from_string(DARK)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4); r=p.add_run(cfg['subtitle']); r.bold=True; r.font.name='Aptos'; r.font.size=Pt(13); r.font.color.rgb=RGBColor.from_string(ACCENT)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(10); r=p.add_run(cfg['period']); r.font.name='Aptos'; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(GRAY)
    t=doc.add_table(rows=1,cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,m in enumerate(cfg['metrics']):
        c=t.cell(0,i); c.width=Inches(1.75); shade(c,'E4F1EC'); margins(c,120); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        a,b=m.split('\n',1); p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(3)
        r=p.add_run(a); r.bold=True; r.font.name='Aptos'; r.font.size=Pt(20); r.font.color.rgb=RGBColor.from_string(ACCENT)
        p=c.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); r=p.add_run(b); r.font.name='Aptos'; r.font.size=Pt(9)
    p=body(doc,cfg['foot'],8.5,7,italic=True); p.runs[0].font.color.rgb=RGBColor.from_string(GRAY)
    box=doc.add_table(rows=1,cols=1); box.alignment=WD_TABLE_ALIGNMENT.CENTER; c=box.cell(0,0); shade(c,'F2F6F4'); margins(c,100); bp=c.paragraphs[0]; bp.paragraph_format.space_after=Pt(0); rr=bp.add_run(cfg['scope']); rr.font.name='Aptos'; rr.font.size=Pt(9.2); rr.font.color.rgb=RGBColor.from_string(DARK)
    p=body(doc,cfg['close'],8.5,7); p.runs[0].font.color.rgb=RGBColor.from_string(GRAY)

def build(lang):
    cfg=LANGS[lang]; doc=Document(cfg['template']); clear_body(doc); first_page(doc,cfg)
    heading(doc,cfg['sections'][0])
    for x in cfg['executive']: body(doc,x,10)
    heading(doc,cfg['sections'][1]); body(doc,cfg['core_intro'],9.2)
    t=doc.add_table(rows=1,cols=6); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    heads={'pt':['Data','País','Local','Fauna / produto','Modalidade','Síntese'],'en':['Date','Country','Location','Wildlife / product','Mode','Summary'],'es':['Fecha','País','Lugar','Fauna / producto','Modalidad','Síntesis']}[lang]
    widths=[0.55,0.78,1.0,1.35,1.2,2.0]
    for i,h in enumerate(heads):
        c=t.cell(0,i); c.width=Inches(widths[i]); shade(c,DARK); margins(c,70); p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(h); r.bold=True; r.font.name='Aptos'; r.font.size=Pt(7.2); r.font.color.rgb=RGBColor(255,255,255)
    for cse in cases:
        row=t.add_row(); no_split(row); desc=cse[lang]; desc=(desc[:170]+'…') if len(desc)>171 else desc
        vals=[cse['date'],val(cse,lang,'country'),cse['location'],val(cse,lang,'fauna'),val(cse,lang,'mode'),desc]
        for i,x in enumerate(vals):
            c=row.cells[i]; c.width=Inches(widths[i]); margins(c,65); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i<2 else WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after=Pt(0); r=p.add_run(x); r.font.name='Aptos'; r.font.size=Pt(6.8)
    heading(doc,cfg['sections'][2])
    for h,x in cfg['analytical']: heading(doc,h,2); body(doc,x,10)
    heading(doc,cfg['sections'][3])
    for i,cse in enumerate(cases,1):
        p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(3)
        r=p.add_run(f"{i:02d}  {val(cse,lang,'country')} — {cse['location']} | {cse['date']}/2026"); r.bold=True; r.font.name='Aptos'; r.font.size=Pt(10.5); r.font.color.rgb=RGBColor.from_string(DARK)
        body(doc,cse[lang],9.5,4); p=doc.add_paragraph(); r=p.add_run(cfg['source_label']); r.bold=True; r.font.name='Aptos'; r.font.size=Pt(8.8); link(p,cse['source'],sources[cse['src']])
    heading(doc,cfg['sections'][4])
    for u in uncertain:
        p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(3)
        r=p.add_run(f"{val(u,lang,'country')} — {u['location']} | {cfg['published_prefix']}{u['published']}"); r.bold=True; r.font.name='Aptos'; r.font.size=Pt(10.5); r.font.color.rgb=RGBColor.from_string(DARK)
        body(doc,u[lang],9.5,4); p=doc.add_paragraph(); r=p.add_run(cfg['source_label']); r.bold=True; r.font.name='Aptos'; r.font.size=Pt(8.8); link(p,u['source'],sources[u['src']])
    heading(doc,cfg['sections'][5])
    for x in cfg['methodology']: body(doc,x,9.5)
    heading(doc,cfg['sections'][6])
    items=[(f"{val(c,lang,'country')} — {c['source']}",sources[c['src']]) for c in cases]+[(f"{val(u,lang,'country')} — {u['source']}",sources[u['src']]) for u in uncertain]
    st=doc.add_table(rows=(len(items)+1)//2,cols=2); st.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,(label,url) in enumerate(items):
        c=st.cell(i//2,i%2); margins(c,70); p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(f'{i+1}. '); r.bold=True; r.font.name='Aptos'; r.font.size=Pt(7.6); link(p,label,url)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(cfg['body_footer']); r.font.name='Aptos'; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)
    for sec in doc.sections:
        if sec.footer.tables:
            c=sec.footer.tables[0].cell(0,0); c.text=cfg['footer']; rr=c.paragraphs[0].runs[0]; rr.font.name='Aptos'; rr.font.size=Pt(8); rr.font.color.rgb=RGBColor.from_string('70817C')
    doc.save(OUT/cfg['outfile'])

for L in ('pt','en','es'): build(L)

idx=ROOT/'reports.json'; current=json.loads(idx.read_text(encoding='utf-8'))
record={'id':'2026-09-08_2026-09-14','period_start':'2026-09-08','period_end':'2026-09-14','year':2026,
'title_pt':'Relatório Semanal de Monitoramento','title_en':'Weekly Monitoring Report','title_es':'Informe Semanal de Monitoreo',
'summary_pt':'Síntese dos casos validados de tráfico e exploração ilegal de fauna registrados entre 8 e 14 de setembro de 2026.',
'summary_en':'Summary of validated wildlife trafficking and illegal exploitation cases recorded between 8 and 14 September 2026.',
'summary_es':'Síntesis de los casos validados de tráfico y explotación ilegal de fauna registrados entre el 8 y el 14 de septiembre de 2026.',
'metrics':{'cases':17,'countries':10,'exact_animals':210,'additional_uncertain_date_records':3},
'pdf_pt':'reports/2026/2026-09-08_2026-09-14/Relatorio_Semanal_Observatorio_Global_08-14_Set_2026_WCS.pdf','docx_pt':'reports/2026/2026-09-08_2026-09-14/Relatorio_Semanal_Observatorio_Global_08-14_Set_2026_WCS.docx',
'pdf_en':'reports/2026/2026-09-08_2026-09-14/Weekly_Report_Global_Observatory_08-14_Sep_2026_WCS.pdf','docx_en':'reports/2026/2026-09-08_2026-09-14/Weekly_Report_Global_Observatory_08-14_Sep_2026_WCS.docx',
'pdf_es':'reports/2026/2026-09-08_2026-09-14/Informe_Semanal_Observatorio_Global_08-14_Sep_2026_WCS.pdf','docx_es':'reports/2026/2026-09-08_2026-09-14/Informe_Semanal_Observatorio_Global_08-14_Sep_2026_WCS.docx','published_at':'2026-09-22'}
current=[r for r in current if r.get('id')!=record['id']]+[record]
current.sort(key=lambda r:r.get('period_end',''),reverse=True)
idx.write_text(json.dumps(current,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
